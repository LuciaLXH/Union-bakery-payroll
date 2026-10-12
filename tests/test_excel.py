import unittest
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from core.excel_handler import create_attendance_template, read_attendance_excel, export_payroll_to_excel, example_inputs
from core.payroll_calc import calculate_payroll, PayrollValidationError


def excel(df):
    buf = BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)
    return buf


class ExcelTests(unittest.TestCase):
    def test_real_headers_without_system_metadata(self):
        df = pd.DataFrame([{"姓名": "甲", "工资": 4500, "绩效": 600, "全勤": 300,
                            "迟到": 7, "加班费": 0, "合计": 5393, "入职时间": "2026-09-01"}])
        inputs = read_attendance_excel(excel(df), default_store="分店甲")
        result = calculate_payroll(inputs).iloc[0]
        self.assertEqual(result["分店"], "分店甲")
        self.assertEqual(result["原表合计"], 5393)
        self.assertEqual(result["合计差额"], 0)

    def test_blank_template_really_has_no_employees(self):
        content = create_attendance_template()
        self.assertEqual(len(pd.read_excel(BytesIO(content))), 0)
        with self.assertRaises(PayrollValidationError):
            read_attendance_excel(BytesIO(content))

    def test_duplicate_excel_headers_are_not_silently_renamed(self):
        df = pd.DataFrame([["甲", 4500, 600, 300, 5000]], columns=["姓名", "工资", "绩效", "全勤", "工资"])
        with self.assertRaisesRegex(ValueError, "重复表头"):
            read_attendance_excel(excel(df))

    def test_example_template_roundtrip(self):
        inputs = read_attendance_excel(BytesIO(create_attendance_template(sample=True)))
        result = calculate_payroll(inputs)
        self.assertEqual(result["合计"].tolist(), [5400, 5220, 4920, 5393, 5450, 5600, 5400, 4740])
        self.assertTrue(result["合计差额"].eq(0).all())

    def test_legacy_file_cannot_silently_infer_salary(self):
        with self.assertRaises(PayrollValidationError):
            read_attendance_excel(excel(pd.DataFrame([{"姓名": "甲", "职位": "后厨", "全勤奖": 400, "绩效奖金": 600}])))

    def test_ignored_columns_are_reported(self):
        df = pd.DataFrame([{"姓名": "甲", "工资": 4500, "绩效": 0, "全勤": 0, "社保扣减额": 1000}])
        result = read_attendance_excel(excel(df))
        self.assertEqual(result.attrs["ignored_columns"], ["社保扣减额"])
        self.assertEqual(calculate_payroll(result).iloc[0]["合计"], 4500)

    def test_bad_rows_not_dropped_and_leading_zero_id_preserved(self):
        df = pd.DataFrame([{"姓名": "甲", "员工编号": "001", "工资": 4500, "绩效": 0, "全勤": 0},
                           {"姓名": None, "员工编号": "002", "工资": 1, "绩效": 0, "全勤": 0}])
        with self.assertRaises(PayrollValidationError) as error:
            read_attendance_excel(excel(df))
        self.assertEqual(error.exception.issues[0]["行号"], 3)
        self.assertEqual(read_attendance_excel(excel(df.iloc[:1])).iloc[0]["员工编号"], "001")

    def test_export_matches_calculation_and_never_writes_user_formulas(self):
        inputs = example_inputs().iloc[:1].copy()
        inputs["姓名"] = '=HYPERLINK("https://example.invalid","test")'
        inputs["工作失误项目"] = "=1+1"
        result = calculate_payroll(inputs)
        content = export_payroll_to_excel(result)
        exported = pd.read_excel(BytesIO(content))
        self.assertEqual(exported.iloc[0]["合计"], 5400)
        wb = load_workbook(BytesIO(content))
        sheet = wb["核算明细"]
        names = {c.value: c.column for c in sheet[1]}
        self.assertEqual(sheet.cell(2, names["姓名"]).data_type, "s")
        self.assertEqual(sheet.cell(2, names["工作失误项目"]).data_type, "s")
        self.assertEqual(sheet.freeze_panes, "C2")

    def test_negative_results_cannot_be_exported(self):
        df = example_inputs().iloc[:1].copy()
        df["旷工天数"] = 30
        with self.assertRaisesRegex(ValueError, "异常"):
            export_payroll_to_excel(calculate_payroll(df))


if __name__ == "__main__":
    unittest.main()
