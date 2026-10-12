"""English demo isolation, calculation parity and actual Streamlit workflows."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
import os
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from core.workspace import employee_frame, entry_frame, EMPLOYEE_DEFAULTS
from portfolio.demo import DemoWorkspace
from portfolio.domain import export_english

APP = str(Path(__file__).resolve().parents[1] / 'portfolio_app.py')


class DemoTests(unittest.TestCase):
    def setUp(self):
        self.repo = DemoWorkspace(date(2026, 10, 12))
        self.store = self.repo.stores()[0]
        self.month = self.repo.months(self.store['id'])[0]

    def inputs(self):
        entries = self.repo.entries(self.month['id'])
        return entry_frame(entries, self.store['name']), {e['id']: e['input_version'] for e in entries}

    def test_independent_sessions_and_copy_boundaries(self):
        other = DemoWorkspace(date(2026, 10, 12))
        self.repo.add_store('My trial store')
        self.assertEqual(len(other.stores()), 2)
        people = self.repo.employees(self.store['id'])
        people[0]['monthly_salary'] = 999
        self.assertEqual(self.repo.employees(self.store['id'])[0]['monthly_salary'], 4500)

    def test_recalculate_preserves_history_and_reuses_same_version(self):
        self.assertEqual(self.repo.current_results(self.month['id'])[0]['net_pay'], 5263)
        frame, versions = self.inputs()
        frame.at[0, '迟到'] = 10
        saved = self.repo.save_inputs(self.month['id'], frame, versions, calculate=True)
        self.assertEqual(saved[0]['net_pay'], 5260)
        history = self.repo.history(store_name=self.store['name'], payroll_month=self.month['payroll_month'])
        self.assertEqual([r['net_pay'] for r in history if r['input_snapshot']['employee_code_snapshot'] == 'DEMO-001'], [5263, 5260])
        frame, versions = self.inputs()
        self.repo.save_inputs(self.month['id'], frame, versions, calculate=True)
        self.assertEqual(len(self.repo.history(store_name=self.store['name'], payroll_month=self.month['payroll_month'])), 4)

    def test_save_only_invalidates_current_result(self):
        frame, versions = self.inputs()
        frame.at[0, '迟到'] = 10
        self.repo.save_inputs(self.month['id'], frame, versions)
        self.assertEqual(len(self.repo.current_results(self.month['id'])), 2)
        with self.assertRaisesRegex(ValueError, 'Inputs changed'):
            self.repo.save_inputs(self.month['id'], frame, versions, calculate=True)

    def test_negative_or_invalid_input_does_not_mutate_records(self):
        frame, versions = self.inputs()
        before = deepcopy(self.repo.data)
        frame.at[0, '普通缺勤天数'] = 0
        frame.at[0, '旷工天数'] = 30
        with self.assertRaisesRegex(ValueError, 'negative'):
            self.repo.save_inputs(self.month['id'], frame, versions, calculate=True)
        self.assertEqual(self.repo.data, before)
        frame.at[0, '普通缺勤天数'] = 1
        with self.assertRaises(ValueError):
            self.repo.save_inputs(self.month['id'], frame, versions, calculate=True)
        self.assertEqual(self.repo.data, before)

    def test_employee_updates_preserve_monthly_snapshot_and_batch_is_atomic(self):
        expected = self.repo.employees(self.store['id'])
        frame = employee_frame(expected)
        frame.at[0, '约定工资'] = 4800
        frame.loc[len(frame)] = dict(EMPLOYEE_DEFAULTS, 员工编号='DEMO-004', 姓名='New Baker', 约定工资=-1)
        with self.assertRaises(ValueError):
            self.repo.save_employees(self.store['id'], frame, expected)
        self.assertEqual(self.repo.employees(self.store['id'])[0]['monthly_salary'], 4500)
        frame.at[3, '约定工资'] = 4000
        self.repo.save_employees(self.store['id'], frame, expected)
        self.assertEqual(self.repo.entries(self.month['id'])[0]['salary_snapshot'], 4500)
        _, added = self.repo.start_month(self.store['id'], self.month['payroll_month'])
        self.assertEqual(added, 1)
        self.assertEqual(self.repo.entries(self.month['id'])[0]['salary_snapshot'], 4500)

    def test_cap_optimistic_check_and_bounds(self):
        self.repo.set_cap(self.month['id'], 10000, None)
        with self.assertRaisesRegex(ValueError, 'changed'):
            self.repo.set_cap(self.month['id'], 12000, None)
        with self.assertRaises(ValueError):
            self.repo.set_cap(self.month['id'], -1, 10000)
        self.assertEqual(self.repo.months(self.store['id'])[0]['payroll_cap_snapshot'], 10000)

    def test_excel_is_english_and_text_not_formulas(self):
        results = self.repo.current_results(self.month['id'])
        results[0]['input_snapshot']['employee_name_snapshot'] = '=1+1'
        w = load_workbook(BytesIO(export_english(results)))
        self.assertEqual(w.sheetnames, ['Payroll detail', 'Read me'])
        rows = list(w['Payroll detail'].values)
        item = dict(zip(rows[0], rows[1]))
        self.assertEqual(item['Net pay'], 5263)
        self.assertEqual(w['Payroll detail']['B2'].data_type, 's')
        self.assertFalse(any('\u4e00' <= c <= '\u9fff' for c in ''.join(rows[0])))


class EnglishPageTests(unittest.TestCase):
    def setUp(self):
        # Even when the owner database URL is present, the demo must never connect.
        with patch.dict(os.environ, {'DATABASE_URL': 'postgresql://must-not-be-used'}), patch('psycopg.connect', side_effect=AssertionError('Demo accessed a database')):
            self.at = AppTest.from_file(APP, default_timeout=20).run()
        self.assertEqual(len(self.at.exception), 0)

    def button(self, label):
        return next(b for b in self.at.button if b.label == label)

    def test_english_ui_recalculate_and_reset(self):
        next(w for w in self.at.number_input if w.label == 'Late minutes').set_value(10).run()
        self.assertEqual(len(self.at.get('download_button')), 0)
        self.button('Save & calculate').click().run()
        self.assertEqual(len(self.at.exception), 0)
        repo = self.at.session_state.portfolio_repo
        month = repo.months(repo.stores()[0]['id'])[0]
        self.assertEqual(repo.current_results(month['id'])[0]['net_pay'], 5260)
        self.assertEqual(len(self.at.get('download_button')), 1)
        self.button('Reset my demo').click().run()
        self.assertEqual(next(w for w in self.at.number_input if w.label == 'Late minutes').value, 7)
        self.assertEqual(len(self.at.exception), 0)

    def test_blank_amount_from_table_is_shown_and_rejected_without_crash(self):
        buffer = next(iter(self.at.session_state.portfolio_months.values()))
        buffer['frame'].at[0, '工资'] = None
        buffer['revision'] += 1
        self.at.run()
        self.assertEqual(len(self.at.exception), 0)
        self.button('Save & calculate').click().run()
        self.assertEqual(len(self.at.exception), 0)
        self.assertTrue(any('cannot be blank' in e.value for e in self.at.error))
        self.assertEqual(len(self.at.get('download_button')), 0)

    def test_visible_text_is_english(self):
        for category in ('title', 'header', 'subheader', 'caption', 'info', 'warning', 'success', 'markdown'):
            for element in getattr(self.at, category):
                self.assertFalse(any('\u4e00' <= c <= '\u9fff' for c in str(element.value)), (category, element.value))
        for category in ('button', 'radio', 'selectbox', 'number_input', 'text_area'):
            for element in getattr(self.at, category):
                self.assertFalse(any('\u4e00' <= c <= '\u9fff' for c in element.label), element.label)


def render_named_page(name):
    from portfolio import app
    getattr(app, name)()


class OtherEnglishPageTests(unittest.TestCase):
    def page(self, name):
        at = AppTest.from_function(render_named_page, args=(name,), default_timeout=20).run()
        self.assertEqual(len(at.exception), 0)
        return at

    def test_store_and_employee_creation_through_forms(self):
        at = self.page('stores')
        at.text_input[0].set_value('Union · Northside')
        next(b for b in at.button if b.label == 'Add store').click().run()
        self.assertEqual(len(at.exception), 0)
        self.assertEqual(len(at.session_state.portfolio_repo.stores()), 3)
        at = self.page('employees')
        next(w for w in at.text_input if w.label == 'Employee ID').set_value('DEMO-004')
        next(w for w in at.text_input if w.label == 'Name').set_value('Morgan Sample')
        next(w for w in at.number_input if w.label == 'Base pay incl. allowance (CNY)').set_value(4000)
        next(b for b in at.button if b.label == 'Save employee').click().run()
        self.assertEqual(len(at.exception), 0)
        self.assertEqual(len(at.session_state.portfolio_repo.employees(1)), 4)

    def test_employee_batch_add_and_edit_callback(self):
        from portfolio.app import employee_table_edit
        at = self.page('employees')
        at.radio[0].set_value('Edit as a table').run()
        store_id, buffer = next(iter(at.session_state.portfolio_employees.items()))
        key = f"en_people_{store_id}_{buffer['generation']}_{buffer['revision']}"
        at.session_state[key] = dict(edited_rows={0: {'约定工资': 4800.0}},
            added_rows=[{'员工编号': 'DEMO-004', '姓名': 'Morgan Sample', '约定工资': 4000.0}], deleted_rows=[])
        with patch('portfolio.app.st.session_state', at.session_state):
            employee_table_edit(key, store_id)
        at.run()
        next(b for b in at.button if b.label == 'Save all employee records').click().run()
        self.assertEqual(len(at.exception), 0)
        people = at.session_state.portfolio_repo.employees(store_id)
        self.assertEqual(len(people), 4)
        self.assertEqual(people[0]['monthly_salary'], 4800)
        self.assertEqual(at.session_state.portfolio_repo.entries(1)[0]['salary_snapshot'], 4500)

    def test_history_and_project_notes_are_english(self):
        for name in ('history', 'about'):
            with self.subTest(page=name):
                at = self.page(name)
                for category in ('header', 'subheader', 'caption', 'info', 'markdown'):
                    for element in getattr(at, category):
                        self.assertFalse(any('\u4e00' <= c <= '\u9fff' for c in str(element.value)))

    def test_limit_is_saved_but_not_shown_as_a_metric(self):
        at = self.page('monthly')
        next(b for b in at.button if b.label == 'Set monthly limit').click().run()
        at.checkbox[0].check()
        next(w for w in at.number_input if w.label == 'Monthly limit (CNY)').set_value(15000)
        next(b for b in at.button if b.label == 'Save monthly limit').click().run()
        self.assertEqual(len(at.exception), 0)
        self.assertTrue(any('limit reached' in w.value for w in at.warning))
        self.assertNotIn('15,000', ' '.join(str(w.value) for w in at.metric))


if __name__ == '__main__':
    unittest.main()
